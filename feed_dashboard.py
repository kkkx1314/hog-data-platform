# -*- coding: utf-8 -*-
"""饲料产量看板 —— 数据层（不依赖 Streamlit，桌面和云端共用）。

四个口径：
  all  全部饲料      → 饲料_猪料各品类整理_图表版.xlsx（8 个品类，猪料合计含其他猪）
  hp   配合饲料      → 饲料_配合饲料_猪料整理_图表版.xlsx
  ns   浓缩饲料      → 饲料_浓缩饲料_猪料整理_图表版.xlsx
  tjj  添加剂预混合饲料 → 饲料_添加剂预混合饲料_猪料整理_图表版.xlsx

数据层只存「月度产量（吨）」，环比/同比在页面里按 Excel 同款公式现算：
  环比 = (本月 / 上月 - 1) × 100
  同比 = (本月 / 上年同月 - 1) × 100
保留两位小数，与四份 Excel 的「月度环比 / 月度同比」逐格一致。

本地优先读桌面的四份 Excel（永远最新）；云端没有桌面，读 data/feed_dashboard.json。
"""
import datetime as dt
import json
import math
import os

HERE = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = os.path.join(HERE, 'data', 'feed_dashboard.json')

EXCEL_PATHS = {
    'all': r'D:\CC\Desktop\饲料_猪料各品类整理_图表版.xlsx',
    'hp': r'D:\CC\Desktop\饲料_配合饲料_猪料整理_图表版.xlsx',
    'ns': r'D:\CC\Desktop\饲料_浓缩饲料_猪料整理_图表版.xlsx',
    'tjj': r'D:\CC\Desktop\饲料_添加剂预混合饲料_猪料整理_图表版.xlsx',
}

SCOPES = [
    {'key': 'all', 'label': '全部饲料', 'short': '全部饲料',
     'desc': '全饲料类型合计（猪料合计含其他猪饲料）'},
    {'key': 'hp', 'label': '配合饲料', 'short': '配合饲料', 'desc': '配合饲料'},
    {'key': 'ns', 'label': '浓缩饲料', 'short': '浓缩饲料', 'desc': '浓缩饲料'},
    {'key': 'tjj', 'label': '添加剂预混合饲料', 'short': '添加剂预混料',
     'desc': '添加剂预混合饲料'},
]

# 选择器里的短标签（4 个口径）
SCOPE_LABELS = {s['key']: s['short'] for s in SCOPES}
LABEL_TO_KEY = {s['short']: s['key'] for s in SCOPES}

# 品类显示顺序：合计在前，分项按层级排
CATEGORY_ORDER = [
    '猪料', '仔猪饲料', '哺乳仔猪饲料', '保育饲料', '母猪饲料',
    '后备母猪饲料', '其他母猪饲料', '生长肥育猪饲料', '其他猪饲料',
]

CATEGORY_LABELS = {
    '猪料': '猪料（合计）',
    '仔猪饲料': '仔猪饲料（=哺乳+保育）',
    '哺乳仔猪饲料': '哺乳仔猪饲料',
    '保育饲料': '保育饲料',
    '母猪饲料': '母猪饲料（=后备+其他母猪）',
    '后备母猪饲料': '后备母猪饲料',
    '其他母猪饲料': '其他母猪饲料',
    '生长肥育猪饲料': '生长肥育猪饲料',
    '其他猪饲料': '其他猪饲料',
}

# 9 个品类口径（配合/浓缩/预混）；全部饲料文件没有单列其他猪
CATS9 = ['猪料', '仔猪饲料', '哺乳仔猪饲料', '保育饲料', '母猪饲料',
         '后备母猪饲料', '其他母猪饲料', '生长肥育猪饲料', '其他猪饲料']

NOTE_MD = """
**口径说明**

- 数据来自 `2026年饲料工业生产情况表(本级汇总).xlsx` 的「云南饲基31表」，为云南省月度饲料产量。
- 品类是嵌套口径，**分项不能相加**：
  - 仔猪饲料 = 哺乳仔猪饲料 + 保育饲料
  - 母猪饲料 = 后备母猪饲料 + 其他母猪饲料
  - 猪料（合计）= 仔猪饲料 + 母猪饲料 + 生长肥育猪饲料 + 其他猪饲料
- 「全部饲料」文件未单列其他猪饲料，但猪料合计中已包含它；按品类看时只能选到 8 项。
- 环比 = (本月 ÷ 上月 − 1) × 100；同比 = (本月 ÷ 上年同月 − 1) × 100，保留两位小数。
- 季度产量 = 该季度 3 个月产量之和；季度环比 = 本季 ÷ 上季 − 1；
  季度同比 = 本季 ÷ 上年同季 − 1。只有 3 个月数据齐全的季度才计算。
- 四个口径互为校验：全部饲料 = 配合饲料 + 浓缩饲料 + 添加剂预混合饲料
  （除 2021-12 的历史修订差异外，逐月、逐品类成立）。
"""


def _num(v):
    """把 Excel 单元格转成 float；空值/文本/NaN 一律 None。"""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if math.isnan(f) or math.isinf(f) else f
    if isinstance(v, str):
        s = v.strip().replace(',', '')
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(v).date()
        except Exception:
            return None
    if isinstance(v, str):
        s = v.strip()
        for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y%m%d', '%Y-%m'):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def extract_scope(path, key, label):
    """读取一份工作簿的「月度数据」→ {levels: {'YYYY-MM': {品类: 吨}}, ...}"""
    import openpyxl

    if not os.path.exists(path):
        raise FileNotFoundError(f'找不到 {path}')
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if '月度数据' not in wb.sheetnames:
            raise ValueError(f'{os.path.basename(path)} 缺少「月度数据」工作表')
        ws = wb['月度数据']
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        raise ValueError(f'{os.path.basename(path)} 的「月度数据」为空')
    hdr = [str(c).strip() if c is not None else '' for c in rows[0]]

    # 表头里的品类列；按 CATEGORY_ORDER 排序，未知列附在后面
    raw_cats = [h for h in hdr[1:] if h]
    cats = [c for c in CATEGORY_ORDER if c in raw_cats]
    cats += [c for c in raw_cats if c not in cats]
    if '猪料' not in cats:
        raise ValueError(f'{os.path.basename(path)} 缺少「猪料」列')

    levels = {}
    for row in rows[1:]:
        d = _as_date(row[0]) if row else None
        if not d:
            continue
        m = f'{d.year:04d}-{d.month:02d}'
        rec = {}
        for i, cat in enumerate(hdr):
            if i == 0 or not cat:
                continue
            v = _num(row[i] if i < len(row) else None)
            if v is not None:
                rec[cat] = round(v, 2)
        if rec:
            levels[m] = rec

    months = sorted(levels)
    latest = None
    for m in reversed(months):
        if '猪料' in levels[m]:
            latest = m
            break
    if latest is None and months:
        latest = months[-1]

    return {
        'key': key,
        'label': label,
        'categories': cats,
        'months': months,
        'levels': levels,
        'latest': latest,
    }


def extract_all(paths=None):
    """读四个口径，返回可直接写 JSON 的完整数据。"""
    paths = paths or EXCEL_PATHS
    scopes = {}
    for spec in SCOPES:
        key = spec['key']
        scopes[key] = extract_scope(paths[key], key, spec['label'])

    latest_candidates = [s['latest'] for s in scopes.values() if s['latest']]
    latest = max(latest_candidates) if latest_candidates else None
    return {
        'meta': {
            'generated_at': dt.datetime.now().strftime('%Y-%m-%d %H:%M'),
            'latest': latest,
            'source': '2026年饲料工业生产情况表(本级汇总).xlsx · 云南饲基31表',
            'source_files': {k: os.path.basename(v) for k, v in paths.items()},
            'unit_level': '吨',
            'unit_rate': '%',
            'note': ('仔猪饲料=哺乳+保育；母猪饲料=后备+其他母猪；'
                     '猪料=仔猪+母猪+生长肥育+其他猪。分项不可相加。'),
        },
        'scope_order': [s['key'] for s in SCOPES],
        'scopes': scopes,
    }


def load_json(path=JSON_PATH):
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def find_source():
    """返回可哈希的 source 描述：本地四份 Excel 齐全就用 Excel，否则用云端 JSON。"""
    present = [s for s in SCOPES if os.path.exists(EXCEL_PATHS[s['key']])]
    if len(present) == len(SCOPES):
        return ('excel', tuple(
            (s['key'], EXCEL_PATHS[s['key']], os.path.getmtime(EXCEL_PATHS[s['key']]))
            for s in SCOPES
        ))
    if present:
        missing = [s['short'] for s in SCOPES if s not in present]
        raise FileNotFoundError(
            f'桌面只找到 {len(present)}/4 份饲料 Excel，缺少：{"、".join(missing)}。'
            f'请确认四份《饲料_*_图表版.xlsx》都在桌面。')
    if os.path.exists(JSON_PATH):
        return ('json', JSON_PATH, os.path.getmtime(JSON_PATH))
    raise FileNotFoundError(
        '找不到饲料数据。本地请确认桌面四份《饲料_*_图表版.xlsx》；'
        '云端请先运行 python build_feed_dashboard_data.py 生成 data/feed_dashboard.json。')


def load(source=None):
    source = source or find_source()
    if source[0] == 'excel':
        paths = {k: p for k, p, _ in source[1]}
        return extract_all(paths)
    return load_json(source[1])


def shift_month(month, delta):
    """'2026-09' + 1 -> '2026-10'；跨年自动处理。"""
    y, m = map(int, month.split('-'))
    idx = y * 12 + (m - 1) + delta
    return f'{idx // 12:04d}-{idx % 12 + 1:02d}'


def shift_quarter(quarter, delta):
    """'2026-Q3' + 1 -> '2026-Q4'；跨年自动处理。"""
    y, n = quarter.split('-Q')
    idx = int(y) * 4 + (int(n) - 1) + delta
    return f'{idx // 4:04d}-Q{idx % 4 + 1}'


def quarter_months(quarter):
    """'2026-Q3' -> ['2026-07', '2026-08', '2026-09']"""
    y, n = quarter.split('-Q')
    y, n = int(y), int(n)
    return [f'{y}-{(n - 1) * 3 + i:02d}' for i in (1, 2, 3)]


def quarter_value(scope, quarter, cat):
    """季度产量 = 3 个月之和；任一月缺失则返回 None（不计算不完整季度）。"""
    vals = [value_at(scope, m, cat) for m in quarter_months(quarter)]
    if any(v is None for v in vals):
        return None
    return round(sum(vals), 2)


def pct_change(cur, prev):
    if cur is None or prev in (None, 0):
        return None
    return round((float(cur) / float(prev) - 1) * 100, 2)


def value_at(scope, month, cat):
    return scope.get('levels', {}).get(month, {}).get(cat)


def period_value(scope, period, cat):
    """period 为 'YYYY-MM' 或 'YYYY-Qn'。"""
    if '-Q' in period:
        return quarter_value(scope, period, cat)
    return value_at(scope, period, cat)


def metric_value(scope, period, cat, metric):
    """metric: 'level' | 'mom' | 'yoy'；返回吨或百分数。

    period 为 'YYYY-MM' 时 mom=环比、yoy=同比；
    为 'YYYY-Qn' 时 mom=季度环比、yoy=季度同比。
    """
    if '-Q' in period:
        cur = quarter_value(scope, period, cat)
        if metric == 'level':
            return cur
        step = -1 if metric == 'mom' else -4
        return pct_change(cur, quarter_value(scope, shift_quarter(period, step), cat))
    cur = value_at(scope, period, cat)
    if metric == 'level':
        return cur
    if metric == 'mom':
        return pct_change(cur, value_at(scope, shift_month(period, -1), cat))
    if metric == 'yoy':
        return pct_change(cur, value_at(scope, shift_month(period, -12), cat))
    raise ValueError(f'未知指标 {metric}')
